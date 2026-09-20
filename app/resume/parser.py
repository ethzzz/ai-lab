"""简历文档解析：docx / pdf / txt / md → 纯文本。

只需提取文字喂给 LLM，不保留排版；扫描版 PDF（无文字层）会得到空文本并报错。
"""
import io

MAX_BYTES = 5 * 1024 * 1024  # 5MB
ALLOWED_EXT = {"txt", "md", "docx", "pdf"}


def _ext(filename: str) -> str:
    name = (filename or "").lower()
    return name.rsplit(".", 1)[-1] if "." in name else ""


def extract_text(filename: str, data: bytes) -> str:
    """按扩展名提取纯文本；非法格式 / 超大 / 空文本抛 ValueError（供上层转 400）。"""
    if not data:
        raise ValueError("文件为空")
    if len(data) > MAX_BYTES:
        raise ValueError("文件超过 5MB 上限")
    ext = _ext(filename)
    if ext not in ALLOWED_EXT:
        raise ValueError("不支持的格式，请上传 .docx / .pdf / .txt / .md")
    if ext in ("txt", "md"):
        text = _decode(data)
    elif ext == "docx":
        text = _from_docx(data)
    else:
        text = _from_pdf(data)
    text = (text or "").strip()
    if not text:
        raise ValueError("未能从文件中提取到文本（可能是扫描版 PDF 或空文档）")
    return text


def _decode(data: bytes) -> str:
    """纯文本多编码容错解码。"""
    for enc in ("utf-8", "gbk", "big5", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _from_docx(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as e:  # 依赖缺失与解析失败区分开
        raise ValueError("服务器缺少 python-docx 依赖，无法解析 .docx") from e
    try:
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    except ValueError:
        raise
    except Exception as e:
        raise ValueError("docx 解析失败，请确认文件为有效的 .docx") from e


def _from_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise ValueError("服务器缺少 pypdf 依赖，无法解析 .pdf") from e
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("PDF 已加密，无法解析")
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except ValueError:
        raise
    except Exception as e:
        raise ValueError("pdf 解析失败，请确认文件为有效的文字版 .pdf") from e
